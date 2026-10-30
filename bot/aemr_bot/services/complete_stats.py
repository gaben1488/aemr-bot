"""Readable XLSX export, built from bounded database batches."""
from __future__ import annotations

import asyncio
from datetime import UTC, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font
from sqlalchemy import func, select, text
from sqlalchemy.orm import selectinload

from aemr_bot.db.models import Appeal, Message
from aemr_bot.services import appeals as appeals_service
from aemr_bot.services import privacy_epoch, sla, stats, work_schedule
from aemr_bot.utils.pii_mask import mask_phone
from aemr_bot.utils.text_pages import split_text_units

_BATCH_SIZE = 100
_SHEET_LIMIT = 1_048_576


def cell_chunks(value: str) -> list[str]:
    """Keep all text, explicitly represent XML-incompatible control characters."""
    safe = "".join(c if c in "\t\n\r" or 0x20 <= ord(c) <= 0xD7FF
                   or 0xE000 <= ord(c) <= 0xFFFD
                   or 0x10000 <= ord(c) <= 0x10FFFF
                   else f"⟦U+{ord(c):04X}⟧" for c in value)
    # Conservative for Excel's cell limit even if non-BMP characters count
    # as two UTF-16 units; allow room for a formula-protection apostrophe.
    return split_text_units(safe, 32000)


def _cell(value):
    if not isinstance(value, str):
        return value
    return "'" + value if value.startswith(("=", "+", "-", "@", "\t", "\r")) else value


def _date(value) -> datetime | str:
    if value is None:
        return ""
    if value.tzinfo is None:
        value = value.replace(tzinfo=UTC)
    # Excel does not support timezone-aware datetimes. The displayed and
    # sortable value is local Kamchatka time, as stated in the summary.
    return value.astimezone(stats.TZ).replace(tzinfo=None)


def _summary_preview(value: str | None, limit: int = 300) -> str:
    """Show an explicitly shortened preview; the complete text is in history."""
    if not value:
        return ""
    return value if len(value) <= limit else value[:limit] + "… [сокращено]"


class _Sheet:
    def __init__(self, workbook, title: str, headers: list[str], widths: list[int]):
        self.workbook, self.title, self.headers, self.widths = workbook, title, headers, widths
        self.number = 0
        self.rows = 0
        self.ws = None
        self.completed: list[tuple] = []
        self._new()

    def _new(self) -> None:
        from openpyxl.utils import get_column_letter

        if self.ws is not None:
            self.completed.append((self.ws, self.rows))
        self.number += 1
        name = self.title if self.number == 1 else f"{self.title} {self.number}"
        self.ws = self.workbook.create_sheet(name[:31])
        self.ws.freeze_panes = "A2"
        for index, width in enumerate(self.widths, 1):
            self.ws.column_dimensions[get_column_letter(index)].width = width
        cells = []
        for title in self.headers:
            cell = WriteOnlyCell(self.ws, value=title)
            cell.font = Font(bold=True)
            cells.append(cell)
        self.ws.append(cells)
        self.rows = 1

    def append(self, values: list) -> None:
        columns = [cell_chunks(value) if isinstance(value, str) else [value] for value in values]
        for part in range(max(len(column) for column in columns)):
            if self.rows >= _SHEET_LIMIT:
                self._new()
            cells = []
            for column in columns:
                value = column[part] if part < len(column) else column[0] if len(column) == 1 else ""
                cell = WriteOnlyCell(self.ws, value=_cell(value))
                if isinstance(value, datetime):
                    cell.number_format = "DD.MM.YYYY HH:MM"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cells.append(cell)
            self.ws.append(cells)
            self.rows += 1

    def finish(self) -> None:
        from openpyxl.utils import get_column_letter

        for sheet, rows in [*self.completed, (self.ws, self.rows)]:
            sheet.auto_filter.ref = f"A1:{get_column_letter(len(self.headers))}{rows}"


def _attachment_description(item) -> str:
    if not isinstance(item, dict):
        return "Вложение: сведения недоступны"
    label = {"image": "Фото", "video": "Видео", "file": "Файл", "audio": "Аудио"}.get(
        str(item.get("type", "")).lower(), "Вложение"
    )
    raw_payload = item.get("payload")
    payload = raw_payload if isinstance(raw_payload, dict) else {}
    name = item.get("filename") or payload.get("filename") or item.get("name") or ""
    return f"{label}: {name}" if name else label


async def build_xlsx_with_epoch(
    session, period: str, *, window: tuple[datetime | None, datetime, str] | None = None
) -> tuple[bytes, str, int, int]:
    if not session.in_transaction() and session.bind.dialect.name == "postgresql":
        # Keep the selected appeals and their histories at one consistent
        # point in time while other operators continue their work.
        await session.execute(text("SET TRANSACTION ISOLATION LEVEL REPEATABLE READ, READ ONLY"))
    snapshot_epoch = await privacy_epoch.read(session)
    start, end, title = window if window is not None else stats.period_window(period)
    workbook = Workbook(write_only=True)
    summary = _Sheet(workbook, "Сводка", ["Показатель", "Значение", "Пояснение"], [55, 25, 80])
    overview = _Sheet(workbook, "Обращения", ["№", "Связано с №", "Подано", "Житель", "Телефон", "Населённый пункт",
                        "Адрес", "Тема", "Краткое содержание", "Статус", "Последнее зафиксированное событие", "Первый ответ периода", "Итоговый ответ", "Закрыто",
                        "Рабочих часов до первого ответа", "Внутренний срок, рабочих часов", "Первый ответ в срок", "Согласие дано",
                        "Согласие отозвано", "Подписка", "Блокировка"],
                      [9, 16, 19, 22, 17, 25, 34, 25, 55, 20, 19, 19, 19, 19, 18, 22, 20, 19, 19, 15, 15])
    history = _Sheet(workbook, "История", ["Обращение №", "Запись", "Дата", "Автор", "Часть", "Текст"],
                     [16, 12, 19, 24, 12, 75])
    files = _Sheet(workbook, "Вложения", ["Обращение №", "Запись", "Дата", "Вложение"], [16, 12, 19, 60])
    consent = _Sheet(workbook, "Согласия", ["Обращение №", "Подтверждено", "Часть", "Текст согласия"],
                     [16, 19, 12, 75])
    periods = _Sheet(workbook, "Сроки", ["Обращение №", "Начало периода", "Первый ответ",
                     "Окончание оценки", "Норматив, рабочих часов", "Просрочено", "Пояснение"],
                     [16, 20, 20, 20, 24, 15, 45])

    def add_entry(appeal_id, record_no, date, author, text, attachments):
        parts = cell_chunks(text or "")
        for i, part in enumerate(parts, 1):
            history.append([appeal_id, record_no, _date(date), author, f"{i}/{len(parts)}", part])
        for item in attachments or []:
            files.append([appeal_id, record_no, _date(date), _attachment_description(item)])

    count, after_id = 0, 0
    first_answered = 0
    response_counts = {"intermediate": 0, "final": 0, "unknown": 0}
    status_counts: dict[str, int] = {}
    while True:
        query = (select(Appeal).options(selectinload(Appeal.user))
                 .where(Appeal.created_at < end, Appeal.id > after_id)
                 .order_by(Appeal.id).limit(_BATCH_SIZE))
        if start is not None:
            query = query.where(Appeal.created_at >= start)
        appeals = list(await session.scalars(query))
        if not appeals:
            break
        after_id = appeals[-1].id
        first_replies = await appeals_service.first_operator_reply_times(session, [a.id for a in appeals])
        latest_rows = await session.execute(
            select(Message.appeal_id, func.max(Message.created_at))
            .where(Message.appeal_id.in_([a.id for a in appeals]))
            .group_by(Message.appeal_id)
        )
        latest_messages = dict(latest_rows.all())
        for appeal in appeals:
            user = appeal.user
            elapsed: float | str = ""
            on_time = ""
            first_reply = first_replies.get(appeal.id)
            if first_reply is not None:
                first_answered += 1
            status_label = stats._status_label(appeal.status)
            status_counts[status_label] = status_counts.get(status_label, 0) + 1
            anchor = appeal.sla_anchor_at or appeal.created_at
            action_dates = [_date(date) for date in (appeal.created_at, appeal.answered_at,
                            appeal.closed_at, appeal.sla_anchor_at, latest_messages.get(appeal.id)) if date]
            latest_action = max(action_dates) if action_dates else ""
            if first_reply and anchor:
                elapsed = round(sla.business_seconds_between(anchor, first_reply, schedule=work_schedule.for_appeal(appeal)) / 3600, 2)
                if appeal.sla_hours_at_start is None or appeal.sla_schedule_at_start is None:
                    on_time = "Нет снимка срока или календаря"
                else:
                    on_time = "нет" if sla.is_overdue(anchor, first_reply, appeal.sla_hours_at_start, schedule=work_schedule.for_appeal(appeal)) else "да"
            overview.append([appeal.id, appeal.related_appeal_id or "", _date(appeal.created_at), getattr(user, "first_name", "") or "",
                             mask_phone(getattr(user, "phone", None)), appeal.locality or "",
                             appeal.address or "", appeal.topic or "", _summary_preview(appeal.summary),
                             stats._status_label(appeal.status), latest_action,
                             _date(first_reply), _date(appeal.answered_at), _date(appeal.closed_at), elapsed,
                             appeal.sla_hours_at_start if appeal.sla_hours_at_start is not None else "Не сохранился", on_time,
                             _date(getattr(user, "consent_pdn_at", None)),
                             _date(getattr(user, "consent_revoked_at", None)),
                             "да" if getattr(user, "subscribed_broadcast", False) else "нет",
                             "да" if getattr(user, "is_blocked", False) else "нет"])
            for period_record in appeal.sla_history or []:
                if not isinstance(period_record, dict):
                    continue
                def period_date(field, record=period_record):
                    try:
                        from datetime import datetime
                        return _date(datetime.fromisoformat(record[field]))
                    except (ValueError, KeyError, TypeError):
                        return ""
                known = (period_record.get("assessment") == "at_reopen"
                         and bool(period_record.get("first_response"))
                         and isinstance(period_record.get("overdue"), bool))
                legacy = period_record.get("assessment") == "at_reopen_legacy"
                periods.append([appeal.id, period_date("start"), period_date("first_response"),
                                period_date("end"), "Не сохранился" if legacy else period_record.get("sla_hours") or "Нет данных",
                                ("да" if period_record["overdue"] else "нет") if known else "Нет данных",
                                "Исходный норматив неизвестен" if legacy else
                                "Нет достоверной оценки первого ответа" if not known else
                                "Оценка рассчитана при возобновлении обращения"])
            records = getattr(user, "consent_records", None) or []
            if not records:
                consent.append([appeal.id, "", "", "Текст ранее данного согласия не сохранён"])
            for record in records:
                if not isinstance(record, dict) or not isinstance(record.get("text"), str):
                    continue
                try:
                    from datetime import datetime
                    accepted = _date(datetime.fromisoformat(record["accepted_at"]))
                except (ValueError, KeyError, TypeError):
                    accepted = "дата не сохранилась"
                parts = cell_chunks(record["text"])
                for i, part in enumerate(parts, 1):
                    consent.append([appeal.id, accepted, f"{i}/{len(parts)}", part])
        messages = await session.stream_scalars(
            select(Message).where(Message.appeal_id.in_([a.id for a in appeals]))
            .order_by(Message.appeal_id, Message.created_at, Message.id)
            .execution_options(yield_per=500)
        )
        index = 0
        record_no = 0
        async for message in messages:
            while index < len(appeals) and appeals[index].id < message.appeal_id:
                a = appeals[index]
                if record_no == 0:
                    add_entry(a.id, 1, a.created_at, "Обращение жителя", a.summary, a.attachments)
                index += 1
                record_no = 0
            if index >= len(appeals):
                break
            appeal = appeals[index]
            if record_no == 0:
                add_entry(appeal.id, 1, appeal.created_at, "Обращение жителя", appeal.summary, appeal.attachments)
                record_no = 1
            record_no += 1
            author = {"from_user": "Житель", "from_operator": "Специалист", "system": "Служебная запись"}.get(
                message.direction, "Запись"
            )
            if message.direction == "from_operator":
                from aemr_bot.services.message_labels import response_label
                author = response_label(message)
                kind = message.response_kind if message.response_kind in ("intermediate", "final") else "unknown"
                response_counts[kind] += 1
            add_entry(appeal.id, record_no, message.created_at, author, message.text, message.attachments)
        while index < len(appeals):
            appeal = appeals[index]
            if record_no == 0:
                add_entry(appeal.id, 1, appeal.created_at, "Обращение жителя", appeal.summary, appeal.attachments)
            index += 1
            record_no = 0
        count += len(appeals)

    summary.append(["Период", title, "Обращения отбираются по дате поступления."])
    summary.append(["Начало периода", _date(start) if start else "Без ограничения", "Включительно; время Камчатки."])
    end_aware = end if end.tzinfo is not None else end.replace(tzinfo=UTC)
    summary.append(["Конец периода", end_aware.astimezone(stats.TZ).isoformat(timespec="microseconds"),
                    "Не включительно; время Камчатки."])
    summary.append(["Обращений", count, "Количество выбранных обращений, а не сообщений."])
    summary.append(["Последнее зафиксированное событие", "", "Последняя сохранённая дата подачи, переписки, ответа, закрытия или начала нового периода; не все действия оператора имеют дату в карточке."])
    for label, total in sorted(status_counts.items()):
        summary.append([f"Статус: {label}", total, "Текущее состояние выбранных обращений."])
    summary.append(["Обращений с первым ответом в текущем периоде", first_answered,
                    "Первый ответ после начала текущего периода работы обращения; не число ответов."])
    summary.append(["Промежуточных ответов (сообщений)", response_counts["intermediate"],
                    "Количество сохранённых сообщений этого типа за всю историю выбранных обращений."])
    summary.append(["Итоговых ответов (сообщений)", response_counts["final"],
                    "Количество сохранённых сообщений этого типа за всю историю выбранных обращений."])
    summary.append(["Ответов без сохранённого типа (сообщений)", response_counts["unknown"],
                    "Исторические ответы без достоверной классификации не отнесены к другим типам."])

    for sheet in (summary, overview, history, files, consent, periods):
        sheet.finish()
    output = BytesIO()
    await asyncio.to_thread(workbook.save, output)
    return output.getvalue(), title, count, snapshot_epoch


async def build_xlsx(session, period: str) -> tuple[bytes, str, int]:
    content, title, count, _ = await build_xlsx_with_epoch(session, period)
    return content, title, count
