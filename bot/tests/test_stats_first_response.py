"""The compact report measures the first reply in its saved response period."""
from datetime import datetime, timedelta
from io import BytesIO
from types import SimpleNamespace
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook

from aemr_bot.services import stats


def appeal(**changes):
    start = datetime(2026, 10, 26, 9, tzinfo=ZoneInfo("Asia/Kamchatka"))
    values = dict(
        id=1, created_at=start, sla_anchor_at=None,
        answered_at=start + timedelta(hours=8), status="answered", user=None,
        locality="", address="", topic="", summary="", sla_hours_at_start=4,
        sla_schedule_at_start={"version": 1, "timezone": "Asia/Kamchatka",
            "start": 9, "end": 18, "lunch_start": 12, "lunch_end": 13, "holidays": []},
        messages=[SimpleNamespace(direction="from_operator", created_at=start + timedelta(hours=1),
                                  text="Промежуточный"),
                  SimpleNamespace(direction="from_operator", created_at=start + timedelta(hours=8),
                                  text="Итоговый")],
    )
    values.update(changes)
    return SimpleNamespace(**values)


def sheet(rows, *args):
    return load_workbook(BytesIO(stats._render_workbook(rows, *args))).active


def test_sla_and_elapsed_stop_at_first_reply_not_final_reply():
    ws = sheet([appeal()])
    assert ws.cell(2, 12).value == 1
    assert ws.cell(2, 13).value == "да"
    assert ws.cell(2, 11).value == "Итоговый"


def test_header_does_not_claim_one_fixed_target_for_mixed_periods():
    ws = sheet([appeal(sla_hours_at_start=2), appeal(id=2, sla_hours_at_start=8)])
    assert "4ч" not in ws.cell(1, 13).value
    assert "период" in ws.cell(1, 13).value


def test_answered_at_does_not_invent_a_missing_first_reply():
    ws = sheet([appeal(messages=[])])
    assert ws.cell(2, 12).value is None
    assert ws.cell(2, 13).value is None


@pytest.mark.parametrize("missing", ["sla_hours_at_start", "sla_schedule_at_start"])
def test_missing_saved_target_or_calendar_is_unknown(missing):
    ws = sheet([appeal(**{missing: None})])
    assert ws.cell(2, 13).value == "Нет снимка срока или календаря"


def test_loaded_messages_from_an_earlier_period_do_not_count():
    row = appeal()
    row.sla_anchor_at = row.created_at + timedelta(hours=2)
    row.messages = row.messages[:1]
    ws = sheet([row])
    assert ws.cell(2, 13).value is None


async def test_database_report_bulk_first_reply_stops_at_intermediate(session):
    from aemr_bot.db.models import Appeal, Message, User
    example = appeal()
    user = User(max_user_id=991502)
    session.add(user)
    await session.flush()
    row = Appeal(user_id=user.id, summary="Текст", status="answered",
        created_at=example.created_at, answered_at=example.answered_at,
        sla_hours_at_start=4, sla_schedule_at_start=example.sla_schedule_at_start)
    session.add(row)
    await session.flush()
    for message in example.messages:
        session.add(Message(appeal_id=row.id, direction=message.direction,
                            text=message.text, created_at=message.created_at))
    await session.flush()
    content, _, count = await stats.build_xlsx(session, "all")
    ws = load_workbook(BytesIO(content)).active
    assert count == 1
    assert ws.cell(2, 12).value == 1
    assert ws.cell(2, 13).value == "да"
