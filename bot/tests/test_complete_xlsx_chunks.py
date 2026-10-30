from aemr_bot.services.complete_stats import cell_chunks
from io import BytesIO
from datetime import datetime

import pytest
from openpyxl import load_workbook


def test_excel_long_text_is_preserved_and_each_chunk_fits():
    original = "Ж" * 70000
    chunks = cell_chunks(original)
    assert "".join(chunks) == original
    assert all(len(chunk) <= 32766 for chunk in chunks)


def test_invalid_xml_character_is_identified_instead_of_dropped():
    assert cell_chunks("начало\x01конец") == ["начало⟦U+0001⟧конец"]


@pytest.mark.asyncio
async def test_report_includes_all_history_and_attachment_labels(session):
    from aemr_bot.db.models import Appeal, Message, User
    from aemr_bot.services.complete_stats import build_xlsx

    user = User(max_user_id=490101, first_name="Иван")
    session.add(user)
    await session.flush()
    appeal = Appeal(user_id=user.id, topic="Дорога", summary="Яма " * 3000,
                    attachments=[{"type": "image", "payload": {"token": "never-export-token"}}])
    session.add(appeal)
    await session.flush()
    session.add_all([Message(appeal_id=appeal.id, direction="from_user", text=f"Дополнение {i}")
                     for i in range(50)])
    await session.flush()
    content, _, count = await build_xlsx(session, "all")
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    rows = list(workbook["История"].values)
    assert count == 1
    assert len(rows) == 52  # header, original, 50 messages
    assert rows[1][-1] == "Яма " * 3000
    assert rows[-1][-1] == "Дополнение 49"
    attachment = list(workbook["Вложения"].values)[1]
    assert "Фото" in attachment[-1]
    assert "never-export-token" not in str(attachment)
    overview = list(workbook["Обращения"].values)
    columns = {name: index for index, name in enumerate(overview[0])}
    assert overview[1][columns["Краткое содержание"]].endswith("[сокращено]")
    assert isinstance(overview[1][columns["Последнее зафиксированное событие"]], datetime)


@pytest.mark.asyncio
async def test_summary_distinguishes_appeals_from_response_messages(session):
    from aemr_bot.db.models import Appeal, Message, User
    from aemr_bot.services.complete_stats import build_xlsx

    user = User(max_user_id=490102, first_name="Анна")
    session.add(user)
    await session.flush()
    appeal = Appeal(user_id=user.id, topic="Дорога", summary="Текст")
    session.add(appeal)
    await session.flush()
    session.add_all([
        Message(appeal_id=appeal.id, direction="from_operator", text="Принято", response_kind="intermediate"),
        Message(appeal_id=appeal.id, direction="from_operator", text="Готово", response_kind="final"),
        Message(appeal_id=appeal.id, direction="from_operator", text="Старый ответ", response_kind=None),
    ])
    await session.flush()
    content, _, _ = await build_xlsx(session, "all")
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    summary = {row[0]: row[1] for row in list(workbook["Сводка"].values)[1:]}
    assert summary["Обращений"] == 1
    assert summary["Обращений с первым ответом в текущем периоде"] == 1
    assert summary["Промежуточных ответов (сообщений)"] == 1
    assert summary["Итоговых ответов (сообщений)"] == 1
    assert summary["Ответов без сохранённого типа (сообщений)"] == 1
    overview = list(workbook["Обращения"].values)
    assert isinstance(overview[1][2], datetime)


@pytest.mark.asyncio
async def test_legacy_period_export_does_not_present_estimate_as_fact(session):
    from aemr_bot.db.models import Appeal, User
    from aemr_bot.services.complete_stats import build_xlsx

    user = User(max_user_id=490103, first_name="Анна")
    session.add(user)
    await session.flush()
    appeal = Appeal(user_id=user.id, topic="Дорога", summary="Текст", attachments=[],
                    sla_history=[{"assessment": "at_reopen_legacy", "sla_hours": 4,
                                  "first_response": "2026-10-01T01:00:00+00:00", "overdue": True}])
    session.add(appeal)
    await session.flush()
    content, _, _ = await build_xlsx(session, "all")
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    rows = list(workbook["Сроки"].values)
    assert rows[1][4] == "Не сохранился"
    assert rows[1][5] == "Нет данных"
