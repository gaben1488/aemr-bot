"""Build the review draft of the regulation from its Markdown source.

Requires python-docx. The source is preserved; this file is a draft until
the administration verifies the legal text, formatting, and approval details.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "docs" / "Регламент_AEMR-bot_MAX_v8.md"
DESTINATION = ROOT / "docs" / "Регламент_AEMR-bot_MAX_v8.docx"


def plain(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = re.sub(r"\\([\\`*_{}\[\]()#+.!|<>-])", r"\1", text)
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    return text.strip()


def cells(line: str) -> list[str]:
    return [plain(item) for item in line.strip().strip("|").split("|")]


def build() -> None:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin, section.right_margin = Cm(3), Cm(1)
    section.top_margin = section.bottom_margin = Cm(2)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(14)
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0
    draft = document.add_paragraph("ПРОЕКТ")
    draft.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    source = re.sub(r"<!--.*?-->", "", SOURCE.read_text(encoding="utf-8"), flags=re.S)
    lines = source.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                row = cells(lines[index])
                if not all(re.fullmatch(r":?-{3,}:?", cell or "") for cell in row):
                    table_lines.append(row)
                index += 1
            if table_lines:
                width = max(map(len, table_lines))
                table = document.add_table(rows=0, cols=width)
                table.style = "Table Grid"
                table.autofit = False
                if width == 6:
                    column_widths = [10, 65, 28, 22, 27, 18]
                elif width == 5:
                    column_widths = [10, 55, 40, 40, 25]
                elif width == 4 and "Команда" in table_lines[0]:
                    column_widths = [10, 54, 65, 41]
                elif width == 4:
                    column_widths = [10, 50, 46, 64]
                else:
                    column_widths = [10, 70, 90]
                for column, millimeters in zip(table.columns, column_widths, strict=True):
                    column.width = Cm(millimeters / 10)
                for row_index, row_values in enumerate(table_lines):
                    row = table.add_row()
                    properties = row._tr.get_or_add_trPr()
                    properties.append(OxmlElement("w:cantSplit"))
                    if row_index < 2:
                        properties.append(OxmlElement("w:tblHeader"))
                    for column, value in enumerate(row_values):
                        cell = row.cells[column]
                        cell.width = Cm(column_widths[column] / 10)
                        cell.text = value
                        for paragraph in cell.paragraphs:
                            paragraph.paragraph_format.first_line_indent = Cm(0)
                            paragraph.paragraph_format.keep_with_next = False
                            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                            for run in paragraph.runs:
                                run.font.size = Pt(11)
            continue
        if line.startswith("#"):
            paragraph = document.add_paragraph(plain(line.lstrip("# ")))
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.keep_with_next = True
            for run in paragraph.runs:
                run.bold = True
        else:
            paragraph = document.add_paragraph(plain(line))
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if re.fullmatch(r"Приложение \d+", line):
                paragraph.paragraph_format.page_break_before = True
        index += 1
    section.different_first_page_header_footer = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    document.save(DESTINATION)
    print(f"Saved draft: {DESTINATION}")


if __name__ == "__main__":
    build()
